"""Сборка резюме в .docx (python-docx). Ввод: тело резюме + агрегированные навыки."""

import io

from docx import Document
from docx.shared import Pt, RGBColor

from app.schemas import ResumeData

RED = RGBColor(0xFF, 0x00, 0x00)


def _add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    p.space_before = Pt(10)


def _add_kv(doc: Document, label: str, value: str, red: bool = False) -> None:
    if not (value or "").strip():
        return
    p = doc.add_paragraph()
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    rv = p.add_run(value.strip())
    if red:
        rv.font.color.rgb = RED


def _add_name_runs(p, data: ResumeData, bold: bool = False, size: Pt | None = None) -> None:
    """ФИО отдельными run'ами: фамилия (предзаполнена) — красная."""
    parts = [(data.last_name, True), (data.first_name, False), (data.middle_name, False)]
    for i, (text, red) in enumerate(parts):
        if not (text or "").strip():
            continue
        if i > 0 and any(x.strip() for x, _ in parts[:i]):
            p.add_run(" ")
        r = p.add_run(text.strip())
        if bold:
            r.bold = True
        if size is not None:
            r.font.size = size
        if red:
            r.font.color.rgb = RED


def build_resume_docx(data: ResumeData, skills: dict) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    if any(x.strip() for x in (data.last_name, data.first_name, data.middle_name)):
        _add_name_runs(title, data, bold=True, size=Pt(18))
    else:
        tr = title.add_run("Резюме")
        tr.bold = True
        tr.font.size = Pt(18)

    pos = skills["position"]["position"]
    _add_kv(doc, "Рекомендуемая должность", pos)
    if skills["position"].get("reasons"):
        _add_kv(doc, "Обоснование", ", ".join(skills["position"]["reasons"]))

    # 1. Основная информация
    _add_heading(doc, "Основная информация")
    p = doc.add_paragraph()
    r1 = p.add_run("ФИО: ")
    r1.bold = True
    _add_name_runs(p, data)
    _add_kv(doc, "Пол", {"male": "мужской", "female": "женский"}.get(data.gender, data.gender))
    _add_kv(doc, "Город", data.city, red=True)
    _add_kv(doc, "Телефон", data.phone, red=True)
    _add_kv(
        doc,
        "Дата рождения",
        " ".join(x for x in (data.birth_day, data.birth_month, data.birth_year) if x.strip()),
        red=True,
    )
    _add_kv(doc, "Гражданство", data.citizenship, red=True)
    _add_kv(doc, "Разрешение на работу", data.work_permit, red=True)

    # 2. Навыки
    _add_heading(doc, "Навыки")
    if skills["hard_skills"]:
        _add_kv(doc, "Hard Skills (языки, фреймворки, БД)", ", ".join(skills["hard_skills"]))
    if skills["tools"]:
        _add_kv(doc, "Tools", ", ".join(skills["tools"]))
    if skills["algorithms"]:
        _add_kv(doc, "Алгоритмы (применяю на практике)", ", ".join(skills["algorithms"]))

    # 3. Ссылки
    links = [x for x in (data.github, data.gitverse, data.gitflic) if x.strip()]
    if links:
        _add_heading(doc, "Ссылки")
        for label, value in [("GitHub", data.github), ("GitVerse", data.gitverse), ("GitFlic", data.gitflic)]:
            _add_kv(doc, label, value)

    # 4. Опыт
    exp_rows = []
    for i in data.internships:
        if not (i.organization or i.city).strip():
            continue
        row = i.organization.strip()
        if i.city.strip():
            row += f", {i.city.strip()}"
        if i.period.strip():
            row += f" — {i.period.strip()}"
        if i.goal.strip():
            row += f". Цель: {i.goal.strip()}"
        exp_rows.append(row)
    lm = data.lab_membership
    if lm is not None and (lm.university or lm.own_project).strip():
        row = "Студенческая научно-исследовательская лаборатория"
        if lm.university.strip():
            row += f", {lm.university.strip()}"
        if lm.department.strip():
            row += f", {lm.department.strip()}"
        if lm.supervisor.strip():
            row += f" (руководитель: {lm.supervisor.strip()})"
        if lm.own_project.strip():
            row += f". Личный проект: {lm.own_project.strip()}"
        others = [p for p in (lm.other_projects or []) if p.strip()]
        if others:
            row += ". Проекты лаборатории: " + "; ".join(others)
        exp_rows.append(row)
    if exp_rows:
        _add_heading(doc, "Опыт")
        for row in exp_rows:
            doc.add_paragraph(row, style="List Bullet")

    # 5. Учебная деятельность (курсовые/дипломные + лабораторные)
    edu_work = []
    for p in data.projects:
        if not (p.name or p.problem).strip():
            continue
        row = p.name.strip()
        if p.problem.strip():
            row += f". Задача: {p.problem.strip()}"
        if p.tech_stack:
            row += " Стек: " + ", ".join(p.tech_stack)
        edu_work.append(row)
    for l in data.labs:
        if not (l.subject or l.work_name).strip():
            continue
        row = f"{l.subject.strip()}: {l.work_name.strip()}".strip(": ")
        if l.goal.strip():
            row += f". Цель: {l.goal.strip()}"
        if l.tech_stack:
            row += " Стек: " + ", ".join(l.tech_stack)
        edu_work.append(row)
    if edu_work:
        _add_heading(doc, "Учебная деятельность (курсовые проекты и лабораторные работы)")
        for row in edu_work:
            doc.add_paragraph(row, style="List Bullet")

    # 6. Хакатоны / олимпиады
    hacks = [h for h in data.hackathons if (h.name or h.organizer).strip()]
    if hacks:
        _add_heading(doc, "Хакатоны и олимпиады")
        for h in hacks:
            row = h.name.strip()
            if h.organizer.strip():
                row += f", {h.organizer.strip()}"
            if h.location.strip():
                row += f", {h.location.strip()}"
            if h.month_year.strip():
                row += f", {h.month_year.strip()}"
            if h.role.strip():
                row += f" (роль: {h.role.strip()})"
            if h.tech_stack:
                row += " Стек: " + ", ".join(h.tech_stack)
            doc.add_paragraph(row, style="List Bullet")

    # 7. Образование (ВУЗ и факультет предзаполнены — красным)
    _add_heading(doc, "Образование")
    edu_parts = [
        (data.university, True),
        (data.faculty, True),
        (data.department, False),
        (data.major, False),
        (data.graduation_year, False),
    ]
    filled = [(t.strip(), red) for t, red in edu_parts if (t or "").strip()]
    if filled:
        p = doc.add_paragraph(style="List Bullet")
        for i, (text, red) in enumerate(filled):
            if i > 0:
                p.add_run(" / ")
            r = p.add_run(text)
            if red:
                r.font.color.rgb = RED
    courses = [c for c in data.courses if (c.name or c.month_year).strip()]
    if courses:
        doc.add_paragraph("Курсы дополнительного образования:")
        for c in courses:
            row = c.name.strip()
            if c.month_year.strip():
                row += f", {c.month_year.strip()}"
            if c.tech_stack:
                row += " Стек: " + ", ".join(c.tech_stack)
            doc.add_paragraph(row, style="List Bullet 2")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
